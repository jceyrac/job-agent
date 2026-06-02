"""
cv_extract.py — Extract plain text from CV/resume files for onboarding.

Supports .docx (python-docx), .pdf (best-effort via pdfplumber or pypdf),
.txt/.md (utf-8 decode).  Falls back gracefully when libraries or files
are missing — never raises, always returns a string (possibly empty).
"""

import os

from paths import DATA_DIR as _DATA_DIR


def extract_cv_text_from_bytes(data: bytes, filename: str) -> str:
    """Extract plain text from an in-memory file.

    Args:
        data: raw file bytes
        filename: original filename (used only for extension detection)

    Returns:
        Extracted text, or "" if the format is unsupported or extraction fails.
    """
    ext = os.path.splitext(filename)[1].lower()

    # ── DOCX ──────────────────────────────────────────────────────────────
    if ext in (".docx", ".doc"):
        try:
            from docx import Document
            from io import BytesIO
            doc = Document(BytesIO(data))
            lines = []
            for para in doc.paragraphs:
                text = para.text.strip()
                if text:
                    lines.append(text)
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        text = cell.text.strip()
                        if text:
                            lines.append(text)
            return "\n".join(lines)
        except Exception:
            return ""

    # ── PDF (best-effort) ─────────────────────────────────────────────────
    if ext == ".pdf":
        # Try pdfplumber first, then pypdf
        try:
            import pdfplumber
            from io import BytesIO
            text_parts = []
            with pdfplumber.open(BytesIO(data)) as pdf:
                for page in pdf.pages:
                    t = page.extract_text()
                    if t:
                        text_parts.append(t)
            return "\n".join(text_parts)
        except Exception:
            pass

        try:
            from pypdf import PdfReader
            from io import BytesIO
            reader = PdfReader(BytesIO(data))
            return "\n".join(
                page.extract_text() or "" for page in reader.pages
            )
        except Exception:
            pass

        return ""  # neither library available

    # ── Plain text / Markdown ─────────────────────────────────────────────
    if ext in (".txt", ".md", ".markdown", ".rst", ""):
        try:
            return data.decode("utf-8")
        except UnicodeDecodeError:
            try:
                return data.decode("latin-1")
            except Exception:
                return ""

    return ""


def extract_cv_text(db) -> str:
    """Extract CV text from the configured cv.master_path.

    Args:
        db: JobStorage instance

    Returns:
        Extracted text, or "" if no path is set or the file doesn't exist.
    """
    path = db.get_config("cv.master_path")
    if not path or not os.path.exists(path):
        return ""
    ext = os.path.splitext(path)[1].lower()
    try:
        with open(path, "rb") as f:
            data = f.read()
        return extract_cv_text_from_bytes(data, f"cv{ext}")
    except Exception:
        return ""


def save_uploaded_cv(data: bytes, filename: str, db) -> str:
    """Save an uploaded CV file to the data/ directory and set cv.master_path.

    Args:
        data: raw file bytes
        filename: original filename
        db: JobStorage instance

    Returns:
        Absolute path the file was saved to.
    """
    ext = os.path.splitext(filename)[1].lower() or ".docx"
    dest = os.path.join(_DATA_DIR, f"cv_master{ext}")
    os.makedirs(_DATA_DIR, exist_ok=True)
    with open(dest, "wb") as f:
        f.write(data)
    db.set_config("cv.master_path", dest)
    return dest
